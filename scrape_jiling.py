"""批量收集 jiling.cc（机灵助手）图像模板提示词。

数据源：https://docs.jiling.cc —— 站方公开的模板教程静态页（VuePress 构建）。
主站 www.jiling.cc/image-templates 是 SPA，模板正文不在公开接口里；
每个模板详情页都指向 docs.jiling.cc/templates/<slug>.html，
提示词以 <pre> 代码块形式公开呈现，因此从 sitemap 批量抓取最稳。

产物：
  output/jiling_templates.json   结构化全量数据（含提示词 + 图片本地路径）
  output/images/<slug>/          每个模板的效果图文件
  output/jiling_templates.xlsx   表格版（需 openpyxl，可选）
  output/jiling_templates.md     阅读版
  output/raw_html/               原始 HTML 备份（便于重解析）

下载图片：默认开启，加 --no-images 可跳过。
"""

from __future__ import annotations

import concurrent.futures as cf
import html
import json
import os
import re
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

BASE = "https://docs.jiling.cc"
SITEMAP = f"{BASE}/sitemap.xml"
ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
RAW = OUT / "raw_html"
IMG_DIR = OUT / "images"

DOWNLOAD_IMAGES = True
MAX_IMAGES_PER_TEMPLATE = 8

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/129.0 Safari/537.36"
)

TAG_RE = re.compile(r"<[^>]+>")
SPACE_RE = re.compile(r"[ \t\u00a0]+")
LINE_BREAK_RE = re.compile(r"<br\s*/?>", re.I)
PRE_RE = re.compile(r"<pre[^>]*>.*?<code[^>]*>(.*?)</code>.*?</pre>", re.S | re.I)
PRE_BARE_RE = re.compile(r"<pre[^>]*>(.*?)</pre>", re.S | re.I)
H2_RE = re.compile(r"<h2[^>]*>(.*?)</h2>(.*?)(?=<h2[ >]|$)", re.S)
STRONG_PAIR_RE = re.compile(
    r"<strong>([^<]{1,40}?)</strong>\s*(?:</[a-z]+>\s*<[^>]+>)*\s*([^<]{1,300}?)\s*(?=<)",
    re.S,
)
LI_RE = re.compile(r"<li[^>]*>(.*?)</li>", re.S)
IMG_RE = re.compile(r'<img[^>]+src="([^"]+)"')
A_HREF_RE = re.compile(r'<a[^>]+href="([^"]+)"')


def fetch(url: str, tries: int = 3, timeout: int = 30) -> str:
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(
                url,
                headers={
                    "User-Agent": UA,
                    "Accept": "text/html,application/xhtml+xml",
                    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.6",
                },
            )
            with urllib.request.urlopen(req, timeout=timeout) as r:
                raw = r.read()
            enc = r.headers.get_content_charset() or "utf-8"
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                return raw.decode("utf-8", "replace")
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last = exc
            time.sleep(1.2 * (i + 1))
    raise RuntimeError(f"fetch failed {url}: {last}")


def clean(fragment: str) -> str:
    text = LINE_BREAK_RE.sub("\n", fragment)
    text = re.sub(r"</p>|</div>|</li>|</h[1-6]>|</tr>", "\n", text, flags=re.I)
    text = TAG_RE.sub("", text)
    text = html.unescape(text)
    lines = [SPACE_RE.sub(" ", ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln).strip()


def fetch_bytes(url: str, tries: int = 3, timeout: int = 45) -> bytes:
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(
                url,
                headers={
                    "User-Agent": UA,
                    "Accept": "image/avif,image/webp,image/*,*/*;q=0.8",
                    "Referer": BASE + "/",
                },
            )
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last = exc
            time.sleep(1.0 * (i + 1))
    raise RuntimeError(f"image fetch failed {url}: {last}")


def guess_ext(url: str, blob: bytes) -> str:
    """优先按魔数判定扩展名，退回 URL 后缀。"""
    head = blob[:16]
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if head.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if head.startswith(b"GIF8"):
        return ".gif"
    if head[:4] == b"RIFF" and blob[8:12] == b"WEBP":
        return ".webp"
    if head.startswith(b"BM"):
        return ".bmp"
    if blob[4:8] in (b"ftyp", b"ftyp".ljust(4, b" ")):
        return ".jpg"
    stem = Path(urllib.parse.urlparse(url).path).suffix.lower()
    if stem in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".avif"}:
        return ".webp" if stem == ".avif" else stem
    return ".img"


def download_images(slug: str, urls: list[str]) -> list[dict]:
    """把某模板的效果图存到 output/images/<slug>/，返回元数据列表。"""
    saved: list[dict] = []
    target = IMG_DIR / slug
    target.mkdir(parents=True, exist_ok=True)
    for n, u in enumerate(urls, start=1):
        try:
            blob = fetch_bytes(u)
        except RuntimeError as exc:
            saved.append({"url": u, "local_path": None, "error": str(exc)[:200]})
            continue
        ext = guess_ext(u, blob)
        base = Path(urllib.parse.urlparse(u).path).stem[:60] or "image"
        fp = target / f"{n:02d}-{base}{ext}"
        fp.write_bytes(blob)
        saved.append(
            {
                "url": u,
                "local_path": str(fp.relative_to(OUT)).replace(os.sep, "/"),
                "bytes": len(blob),
                "format": ext.lstrip("."),
            }
        )
    return saved


def split_sections(page: str) -> dict[str, str]:
    """按 <h2> 切分正文，返回 {小标题: 正文}。"""
    body_start = page.find("<main")
    if body_start == -1:
        body_start = page.find("<h1")
    body = page[body_start:] if body_start != -1 else page
    out: dict[str, str] = {}
    for m in H2_RE.finditer(body):
        title = clean(m.group(1))
        out[title] = clean(m.group(2))
    return out


def extract_prompt(page: str, sections: dict[str, str]) -> tuple[str, str]:
    """返回 (提示词正文, 提取方式标记)。"""
    # 1) 优先取「提示词怎么改」小节里的 code/pre 块
    for key in ("提示词怎么改", "提示词", "模板提示词"):
        idx = page.find(html.escape(key))
        if idx == -1:
            idx = re.search(key, page)
            idx = idx.start() if idx else -1
        if idx != -1:
            nxt = page.find("<h2", idx)
            window = page[idx : nxt if nxt != -1 else idx + 20000]
            blocks = PRE_RE.findall(window) + PRE_BARE_RE.findall(window)
            if blocks:
                return clean(blocks[0]).strip(), f"section:{key}"
    # 2) 退化：整页第一个 pre/code 块
    blocks = PRE_RE.findall(page) + PRE_BARE_RE.findall(page)
    if blocks:
        return clean(blocks[0]).strip(), "first-pre"
    return "", "not-found"


META_LABELS = [
    "推荐模型",
    "默认画幅",
    "模板分类",
    "默认尺寸",
    "适用场景",
    "生成耗时",
    "参考图",
]


def extract_meta(page: str, sections: dict[str, str]) -> dict[str, str]:
    meta: dict[str, str] = {}
    blob = "\n".join(
        sections.get(k, "") for k in sections if "模板功能" in k or "适用场景" in k
    )
    if not blob:
        blob = clean(page)
    for label in META_LABELS:
        m = re.search(rf"{label}\s*[:：]?\s*(.+)", blob)
        if m:
            meta[label] = SPACE_RE.sub(" ", m.group(1)).strip(" :：").strip()
    return meta


def parse_page(url: str, slug: str) -> dict:
    page = fetch(url)
    (RAW / f"{slug}.html").write_text(page, encoding="utf-8")

    title = ""
    m = re.search(r"<h1[^>]*>(.*?)</h1>", page, re.S)
    if m:
        title = clean(m.group(1))
    if not title:
        m = re.search(r"<title>(.*?)</title>", page, re.S)
        title = clean(m.group(1)) if m else slug

    desc = ""
    m = re.search(r"</h1>\s*<p[^>]*>(.*?)</p>", page, re.S)
    if m:
        desc = clean(m.group(1)).replace("\n", " ")
    if not desc:
        m = re.search(r'<meta\s+name="description"\s+content="([^"]*)"', page)
        if m:
            raw = html.unescape(m.group(1)).strip()
            # 去掉 SEO 前缀，保留真正的模板描述
            for marker in ("适合", "这个模板", "可查看"):
                idx = raw.find(marker)
                if idx > 0:
                    raw = raw[idx:]
                    break
            desc = raw

    sections = split_sections(page)
    prompt, how = extract_prompt(page, sections)
    meta = extract_meta(page, sections)

    tags: list[str] = []
    tag_block = None
    for key in sections:
        if "模板标签" in key or key.strip() == "标签":
            idx = page.find(key)
            nxt = page.find("<h2", idx) if idx != -1 else -1
            if idx != -1:
                tag_block = page[idx : nxt if nxt != -1 else idx + 4000]
            break
    if tag_block:
        tags = [
            clean(a)
            for a in re.findall(r"<a[^>]*>(.*?)</a>", tag_block, re.S)
            if clean(a) and len(clean(a)) < 24
        ]
        tags = [t for t in dict.fromkeys(tags) if not t.startswith("查看全部")]
    if not tags:
        tm = re.search(
            r'<script[^>]*type="application/ld\+json"[^>]*>(.*?)</script>',
            page,
            re.S,
        )
        if tm:
            try:
                data = json.loads(tm.group(1))
                kw = data.get("keywords") or []
                tags = kw if isinstance(kw, list) else [kw]
            except Exception:  # noqa: BLE001
                pass

    imgs: list[dict] = []
    seen_src: set[str] = set()
    for tag in re.findall(r"<img[^>]+>", page):
        src_m = re.search(r'src="([^"]+)"', tag)
        if not src_m:
            continue
        src = html.unescape(src_m.group(1))
        alt_m = re.search(r'alt="([^"]*)"', tag)
        alt = html.unescape(alt_m.group(1)) if alt_m else ""
        if not src.startswith("http"):
            continue
        low = src.lower()
        if "miniapp" in low or "logo" in low or "favicon" in low:
            continue
        if src in seen_src:
            continue
        seen_src.add(src)
        role = "preview"
        if "效果图" in alt:
            role = "preview"
        elif "示例" in alt or "case" in low:
            role = "example"
        elif "参考" in alt:
            role = "reference"
        imgs.append({"url": src, "alt": alt, "role": role})

    live = next(
        (
            h
            for h in A_HREF_RE.findall(page)
            if "jiling.cc/image-templates/" in h and not h.endswith(".html")
        ),
        f"https://www.jiling.cc/image-templates/{slug}",
    )

    score = ""
    m = re.search(r"(\d\.\d)\s*/\s*5", page) or re.search(
        r'"ratingValue"\s*:\s*"?([\d.]+)', page
    )
    if m:
        score = m.group(1)
    reviews = ""
    m = re.search(r"([\d,]{3,})\s*(?:条评价|人评价|reviews)", page)
    if m:
        reviews = m.group(1)

    steps = sections.get("3步上手", "") or sections.get("三步上手", "")
    tips = sections.get("使用建议", "")
    ref_img = sections.get("参考图建议", "")

    local_images: list[dict] = []
    if DOWNLOAD_IMAGES and imgs:
        local_images = download_images(slug, [i["url"] for i in imgs[:MAX_IMAGES_PER_TEMPLATE]])
    for rec, loc in zip(imgs, local_images):
        rec.update(loc)

    return {
        "slug": slug,
        "title": title,
        "description": desc,
        "category": meta.get("模板分类", ""),
        "model": meta.get("推荐模型", ""),
        "aspect_ratio": meta.get("默认画幅", ""),
        "size": meta.get("默认尺寸", ""),
        "scenarios": meta.get("适用场景", ""),
        "tags": tags,
        "prompt": prompt,
        "prompt_extracted_by": how,
        "steps": steps,
        "reference_image_advice": ref_img,
        "usage_tips": tips,
        "score": score,
        "reviews": reviews,
        "images": imgs,
        "image_count": len([i for i in imgs if i.get("local_path")]),
        "preview_images": [i["url"] for i in imgs],
        "docs_url": url,
        "live_url": live,
    }


def main() -> int:
    global DOWNLOAD_IMAGES  # noqa: PLW0603

    argv = sys.argv[1:]
    if "--no-images" in argv:
        DOWNLOAD_IMAGES = False
        argv.remove("--no-images")
    limit = None
    if "--limit" in argv:
        limit = int(argv[argv.index("--limit") + 1])
    workers = 6
    if "--workers" in argv:
        workers = int(argv[argv.index("--workers") + 1])

    OUT.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(parents=True, exist_ok=True)
    IMG_DIR.mkdir(parents=True, exist_ok=True)

    xml = fetch(SITEMAP)
    urls = re.findall(r"<loc>([^<]+)</loc>", xml)
    tpl_urls = sorted({u for u in urls if "/templates/" in u})
    print(
        f"sitemap 总 URL {len(urls)}，模板页 {len(tpl_urls)}，"
        f"下载图片={'是' if DOWNLOAD_IMAGES else '否'}",
        flush=True,
    )

    if limit:
        tpl_urls = tpl_urls[:limit]
        print(f"[试跑模式] 仅抓前 {limit} 个", flush=True)

    results: list[dict] = []
    errors: list[tuple[str, str]] = []
    t0 = time.time()
    with cf.ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {}
        for u in tpl_urls:
            slug = u.rstrip("/").split("/")[-1].removesuffix(".html")
            futs[ex.submit(parse_page, u, slug)] = (u, slug)
        done = 0
        for fut in cf.as_completed(futs):
            u, slug = futs[fut]
            try:
                results.append(fut.result())
            except Exception as exc:  # noqa: BLE001
                errors.append((u, str(exc)[:300]))
            done += 1
            if done % 25 == 0 or done == len(futs):
                print(f"  {done}/{len(futs)}  elapsed={time.time()-t0:.0f}s", flush=True)

    results.sort(key=lambda r: r["slug"])
    payload = {
        "source": "https://www.jiling.cc/image-templates",
        "scraped_from": BASE,
        "scraped_at": time.strftime("%Y-%m-%d %H:%M:%S %z"),
        "total": len(results),
        "templates": results,
    }
    (OUT / "jiling_templates.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    no_prompt = [r["slug"] for r in results if not r["prompt"]]
    n_imgs = sum(r["image_count"] for r in results)
    bytes_imgs = sum(
        i.get("bytes", 0) for r in results for i in r["images"] if i.get("bytes")
    )
    print(
        f"\n成功 {len(results)} / {len(tpl_urls)}；错误 {len(errors)}；"
        f"无提示词 {len(no_prompt)}；图片 {n_imgs} 张 / {bytes_imgs/1e6:.1f} MB"
    )
    if no_prompt:
        print("  缺提示词:", ", ".join(no_prompt[:20]))
    failed_imgs = [
        (r["slug"], i["url"], i.get("error"))
        for r in results
        for i in r["images"]
        if i.get("local_path") is None and i.get("error")
    ]
    if failed_imgs:
        print(f"  图片下载失败 {len(failed_imgs)} 张，示例:", failed_imgs[0])
    for u, e in errors[:10]:
        print("  ERR", u, e)

    build_markdown(results, OUT / "jiling_templates.md")
    try:
        build_xlsx(results, OUT / "jiling_templates.xlsx")
        print("  已写 xlsx")
    except Exception as e:  # noqa: BLE001
        print("  xlsx 跳过:", e)
    return 0


def build_markdown(results: list[dict], path: Path) -> None:
    parts = [
        "# 机灵助手 (jiling.cc) 图像模板提示词合集",
        "",
        f"来源: https://www.jiling.cc/image-templates ｜ 模板数: {len(results)} ｜ "
        f"抓取时间: {time.strftime('%Y-%m-%d %H:%M')}",
        "",
    ]
    parts.append("## 目录\n")
    for r in results:
        parts.append(f"- [{r['title']}](#{r['slug']})（{r['category']}）")
    parts.append("")
    for r in results:
        parts.append(f"\n---\n\n## {r['title']}  <!-- {r['slug']} -->\n")
        if r["description"]:
            parts.append(f"{r['description']}\n")
        info = " ｜ ".join(
            x
            for x in [
                f"分类: {r['category']}" if r["category"] else "",
                f"模型: {r['model']}" if r["model"] else "",
                f"画幅: {r['aspect_ratio']}" if r["aspect_ratio"] else "",
                f"尺寸: {r['size']}" if r["size"] else "",
                f"评分: {r['score']}/5" if r["score"] else "",
                f"{r['reviews']} 条评价" if r["reviews"] else "",
            ]
            if x
        )
        if info:
            parts.append(f"{info}\n")
        if r["tags"]:
            parts.append("标签: " + " ".join(f"`{t}`" for t in r["tags"]) + "\n")
        parts.append("### 提示词\n")
        parts.append("```text\n" + (r["prompt"] or "（未提取到）") + "\n```\n")
        if r["images"]:
            parts.append("### 效果图\n")
            for i in r["images"]:
                if i.get("local_path"):
                    parts.append(f"![{i['alt'] or r['title']}]({i['local_path']})")
                elif i.get("url"):
                    parts.append(f"- 图片下载失败: {i['url']}")
            parts.append("")
        if r["steps"]:
            parts.append(f"### 使用步骤\n\n{r['steps']}\n")
        if r["usage_tips"]:
            parts.append(f"### 使用建议\n\n{r['usage_tips']}\n")
        parts.append(f"详情: {r['docs_url']} ｜ 在线使用: {r['live_url']}\n")
    path.write_text("\n".join(parts), encoding="utf-8")


def build_xlsx(results: list[dict], path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "templates"
    headers = [
        ("slug", 26),
        ("标题", 26),
        ("分类", 16),
        ("标签", 30),
        ("推荐模型", 22),
        ("画幅", 10),
        ("默认尺寸", 14),
        ("适用场景", 26),
        ("简介", 46),
        ("提示词", 78),
        ("评分", 8),
        ("评价数", 10),
        ("效果图数", 10),
        ("本地图片", 52),
        ("文档页", 40),
        ("在线使用", 40),
    ]
    ws.append([h for h, _ in headers])
    for c, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    for r in results:
        local_imgs = "; ".join(
            i["local_path"] for i in r["images"] if i.get("local_path")
        )
        ws.append(
            [
                r["slug"],
                r["title"],
                r["category"],
                ", ".join(r["tags"]),
                r["model"],
                r["aspect_ratio"],
                r["size"],
                r["scenarios"],
                r["description"],
                r["prompt"],
                r["score"],
                r["reviews"],
                r["image_count"],
                local_imgs,
                r["docs_url"],
                r["live_url"],
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws2 = wb.create_sheet("images")
    ws2.append(["slug", "标题", "序号", "role", "本地路径", "尺寸(bytes)", "格式", "原始URL"])
    for c, w in enumerate([26, 26, 8, 12, 52, 14, 10, 60], start=1):
        ws2.column_dimensions[get_column_letter(c)].width = w
        ws2.cell(row=1, column=c).font = Font(bold=True)
    ws2.freeze_panes = "A2"
    for r in results:
        for n, i in enumerate(r["images"], start=1):
            ws2.append(
                [
                    r["slug"],
                    r["title"],
                    n,
                    i.get("role", ""),
                    i.get("local_path") or "",
                    i.get("bytes") or "",
                    i.get("format") or "",
                    i.get("url", ""),
                ]
            )
    wb.save(path)



if __name__ == "__main__":
    sys.exit(main())
